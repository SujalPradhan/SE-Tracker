"""Excel export matching the legacy per-team workbook format exactly.

Instructors receive these files, so sheet names and header rows must not
drift — note the en-dashes (–) in the score headers and the literal
"Avg Score (12 Calls)" kept for compatibility with the legacy sheets.
"""

import io
import zipfile

from openpyxl import Workbook

import metrics
import services

TEAM_INFO_HEADER = ["Team ID", "Member Name", "Role", "Email"]

CLIENT_CALLS_HEADER = [
    "Call #",
    "Week",
    "Milestone #",
    "Team ID",
    "Member Name",
    "Role",
    "Present",
    "Spoke",
    "Tasks Assigned (Prev Call)",
    "Progress Since Last Call",
    "Evidence Link",
    "Role Fulfillment (0–3)",
    "Client Interaction (0–2)",
    "Initiative (0–2)",
    "Total Score",
    "TA Remarks",
]

MILESTONE_SUBMISSIONS_HEADER = [
    "Milestone #",
    "Team ID",
    "Member Name",
    "Role",
    "Report Section Contributed",
    "Code (Y/N)",
    "Tests (Y/N)",
    "Demo (Y/N)",
    "Evidence Link",
    "Remarks",
]

CALL_SUMMARY_HEADER = [
    "Team ID",
    "Member Name",
    "Role",
    "Calls Attended",
    "Avg Score (12 Calls)",
    "Calls With Evidence",
    "Silent Calls",
    "Contribution Status",
]

RISK_FLAGS_HEADER = ["Team ID", "Member Name", "Risk Signal", "Evidence"]


def team_filename(team, cohort):
    return f"{cohort}-Team-{team.team_number:03d}.xlsx"


def _yn(flag):
    return "Y" if flag else "N"


def build_team_workbook(team):
    team_id = team.display_number
    wb = Workbook()

    ws = wb.active
    ws.title = "Team_Info"
    ws.append(TEAM_INFO_HEADER)
    for member in team.members:
        ws.append([team_id, member.name, member.role, member.email])

    ws = wb.create_sheet("Client_Calls")
    ws.append(CLIENT_CALLS_HEADER)
    for call in services.sorted_calls(team):
        entries = {e.member_id: e for e in call.entries}
        for member in team.members:  # roster order within each call
            entry = entries.get(member.id)
            if not entry:
                continue
            ws.append(
                [
                    call.call_no,
                    call.week,
                    call.milestone,
                    team_id,
                    member.name,
                    member.role,
                    _yn(entry.present),
                    _yn(entry.spoke),
                    entry.tasks_assigned_prev,
                    entry.progress,
                    entry.evidence_link,
                    entry.role_fulfillment,  # None stays a blank cell
                    entry.client_interaction,
                    entry.initiative,
                    metrics.call_total(entry.scores),  # always numeric
                    entry.remarks,
                ]
            )

    ws = wb.create_sheet("Milestone_Submissions")
    ws.append(MILESTONE_SUBMISSIONS_HEADER)
    subs = sorted(team.submissions, key=lambda s: (s.milestone, s.member_id))
    for sub in subs:
        ws.append(
            [
                sub.milestone,
                team_id,
                sub.member.name,
                sub.member.role,
                sub.report_section,
                sub.code,
                sub.tests,
                sub.demo,
                sub.evidence_link,
                sub.remarks,
            ]
        )

    summary = services.team_summary(team)

    ws = wb.create_sheet("Call_Summary")
    ws.append(CALL_SUMMARY_HEADER)
    for member, m, _dot in summary:
        avg = m["avg_score"]
        ws.append(
            [
                team_id,
                member.name,
                member.role,
                m["calls_attended"],
                "—" if avg is None else round(avg, 2),
                m["calls_with_evidence"],
                m["silent_calls"],
                m["contribution_status"],
            ]
        )

    ws = wb.create_sheet("Risk_Flags")
    ws.append(RISK_FLAGS_HEADER)
    for member, m, _dot in summary:
        ws.append([team_id, member.name, m["risk_signal"], m["risk_evidence"]])

    return wb


def workbook_bytes(team):
    buf = io.BytesIO()
    build_team_workbook(team).save(buf)
    return buf.getvalue()


def all_teams_zip(teams, cohort):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for team in teams:
            zf.writestr(team_filename(team, cohort), workbook_bytes(team))
    buf.seek(0)
    return buf
