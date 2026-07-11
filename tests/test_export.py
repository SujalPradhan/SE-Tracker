"""The legacy workbook format: sheet names and headers must match byte-for-byte.

Header literals here are written out in full (including the en-dashes – in the
score columns) on purpose — they must NOT be imported from export.py, or the
test would pass even if the app's headers drifted from the legacy format.
"""

import io
import zipfile

from openpyxl import load_workbook

import export
import services
from conftest import make_call, make_ta, make_team
from models import MilestoneSubmission, db

EXPECTED_SHEETS = [
    "Team_Info",
    "Client_Calls",
    "Milestone_Submissions",
    "Call_Summary",
    "Risk_Flags",
]

EXPECTED_HEADERS = {
    "Team_Info": ("Team ID", "Member Name", "Role", "Email"),
    "Client_Calls": (
        "Call #",
        "Week",
        "Milestone #",
        "Team ID",
        "Member Name",
        "Role",
        "Present",
        "Spoke",
        "Tasks Assigned (Prev Call)",
        "Progress Since Last Call",
        "Evidence Link",
        "Role Fulfillment (0–3)",
        "Client Interaction (0–2)",
        "Initiative (0–2)",
        "Total Score",
        "TA Remarks",
    ),
    "Milestone_Submissions": (
        "Milestone #",
        "Team ID",
        "Member Name",
        "Role",
        "Report Section Contributed",
        "Code (Y/N)",
        "Tests (Y/N)",
        "Demo (Y/N)",
        "Evidence Link",
        "Remarks",
    ),
    "Call_Summary": (
        "Team ID",
        "Member Name",
        "Role",
        "Calls Attended",
        "Avg Score (12 Calls)",  # legacy header kept verbatim
        "Calls With Evidence",
        "Silent Calls",
        "Contribution Status",
    ),
    "Risk_Flags": ("Team ID", "Member Name", "Risk Signal", "Evidence"),
}


def exported_team(app):
    team = make_team(number=96)
    make_call(
        team,
        1,
        milestone="M1",
        week="W1",
        entries=[
            {
                "present": True,
                "spoke": True,
                "tasks_assigned_prev": "",
                "progress": "Set up repo",
                "evidence_link": "https://github.com/org/repo",
                "role_fulfillment": 2.5,
                "client_interaction": None,  # blank cell must stay blank
                "initiative": 1.0,
                "remarks": "Good",
            },
            {"present": False, "spoke": False},
        ],
    )
    db.session.add(
        MilestoneSubmission(
            team_id=team.id,
            member_id=team.members[0].id,
            milestone="M1",
            report_section="Intro",
            code="Y",
            tests="N",
            demo="",
            evidence_link="https://github.com/org/repo",
            remarks="On time",
        )
    )
    db.session.commit()
    return team


def load(team):
    return load_workbook(io.BytesIO(export.workbook_bytes(team)))


def test_sheet_names_exact(app):
    wb = load(exported_team(app))
    assert wb.sheetnames == EXPECTED_SHEETS


def test_header_rows_byte_for_byte(app):
    wb = load(exported_team(app))
    for sheet, expected in EXPECTED_HEADERS.items():
        header = next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True))
        assert header == expected, f"{sheet} header drifted from the legacy format"


def test_client_calls_row_values(app):
    team = exported_team(app)
    wb = load(team)
    rows = list(wb["Client_Calls"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2  # one row per member for the single call

    first = rows[0]
    assert first[0] == 1                      # Call #
    assert first[3] == "096"                  # Team ID zero-padded
    assert first[6] == "Y" and first[7] == "Y"  # Present / Spoke as Y/N
    assert first[11] == 2.5                   # 0.5-step score preserved
    assert first[12] is None                  # blank score cell stays blank
    assert first[14] == 3.5                   # Total Score always numeric

    absent = rows[1]
    assert absent[6] == "N" and absent[7] == "N"
    assert absent[14] == 0                    # unscored call still numeric


def test_summary_and_risk_sheets_are_plain_values(app):
    team = exported_team(app)
    wb = load(team)

    summary = list(wb["Call_Summary"].iter_rows(min_row=2, values_only=True))
    assert len(summary) == len(team.members)
    member1 = summary[0]
    assert member1[0] == "096"
    assert member1[3] == 1                    # calls attended
    assert member1[4] == 3.5                  # avg over 1 call, plain value
    assert member1[7] == "🟡 At Risk"         # emoji kept in status cells

    risk = list(wb["Risk_Flags"].iter_rows(min_row=2, values_only=True))
    assert len(risk) == len(team.members)
    assert risk[0][2] in ("🟢 No Risk", "🔴 Risk: Low Contribution")


def test_avg_shows_dash_with_zero_calls(app):
    team = make_team(number=7)  # no calls logged at all
    wb = load(team)
    row = next(wb["Call_Summary"].iter_rows(min_row=2, values_only=True))
    assert row[4] == "—"
    assert row[7] == "— No calls yet"


def test_filename_zero_pads_team_number(app):
    team = make_team(number=96)
    assert export.team_filename(team, "MAY2026") == "MAY2026-Team-096.xlsx"
    team2 = make_team(number=104)
    assert export.team_filename(team2, "MAY2026") == "MAY2026-Team-104.xlsx"


def test_export_all_zip_contains_every_team(app):
    ta = make_ta()
    make_team(number=96, ta=ta)
    make_team(number=104, ta=ta)
    buf = export.all_teams_zip(services.teams_for(ta), "MAY2026")
    with zipfile.ZipFile(buf) as zf:
        assert sorted(zf.namelist()) == [
            "MAY2026-Team-096.xlsx",
            "MAY2026-Team-104.xlsx",
        ]
        # Each entry is a valid workbook with the exact sheet list.
        wb = load_workbook(io.BytesIO(zf.read("MAY2026-Team-096.xlsx")))
        assert wb.sheetnames == EXPECTED_SHEETS
